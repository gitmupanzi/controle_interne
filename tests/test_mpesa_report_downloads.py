"""Les colonnes techniques restent internes aux contrôles, hors rapport opérationnel."""
from io import BytesIO
import os
from pathlib import Path
import unittest

import pandas as pd
from docx import Document

from credit_app.services.mpesa_analysis import (
    G2_CLASSIFIED_TRANSACTION_COLUMNS,
    MpesaPreparedData,
    build_g2_daily_savings_report,
    create_excel_export,
    create_g2_dat_word,
    prepare_transactions,
)


class ReportDownloadTests(unittest.TestCase):
    def setUp(self):
        self.detail = pd.DataFrame([{
            "date": pd.Timestamp("2026-09-11 12:00:00"), "receipt_no": "TEST-EXPORT",
            "currency_code": "CDF", "details_rapport": "Depot normal", "sens_flux": "Entree",
            "opposite_party": "Client test", "montant": 1250.5, "montant_entree": 1250.5,
            "montant_sortie": 0, "balance_numeric": 2000, "duree": 0, "compte_cree": "Oui",
            "statut_rapprochement": "Rapproche exact", "incluse_synthese": True,
        }])

    def test_word_hides_control_but_preserves_data(self):
        original = self.detail.copy(deep=True)
        content = create_g2_dat_word(
            {"rapport_journalier_detail": self.detail, "analysis_source_label": "Solution Numérique"},
            period_text="11/09/2026", direction_label="Tous",
            generated_at=pd.Timestamp("2026-09-12"),
        )
        document = Document(BytesIO(content))
        tables = [t for t in document.tables if t.rows and t.rows[0].cells[0].text == "date"]
        self.assertEqual(len(tables), 1)
        self.assertEqual([c.text for c in tables[0].rows[0].cells],
                         [c for c in G2_CLASSIFIED_TRANSACTION_COLUMNS if c != "statut_rapprochement"])
        text = " ".join(c.text for row in tables[0].rows for c in row.cells)
        self.assertNotIn("Controle Turbo/G2", text)
        self.assertIn("TEST-EXPORT", text)
        pd.testing.assert_frame_equal(self.detail, original)

    def test_excel_hides_control_only_in_operational_details(self):
        original = self.detail.copy(deep=True)
        for renamed in (False, True):
            content = create_excel_export({
                "rapport_journalier_detail": self.detail,
                "rapport_g2_detail": self.detail,
                "rapport_journalier_anomalies": self.detail,
            }, rename_user_columns=renamed)
            with pd.ExcelFile(BytesIO(content)) as workbook:
                self.assertEqual(workbook.sheet_names, ["Rapport_G2_Detail", "Rapport_Journalier_Detail", "Anomalies_G2"])
                audit = pd.read_excel(workbook, sheet_name="Anomalies_G2")
                for sheet in ("Rapport_G2_Detail", "Rapport_Journalier_Detail"):
                    detail = pd.read_excel(workbook, sheet_name=sheet)
                    self.assertEqual(len(detail), 1)
                    self.assertEqual(len(audit.columns), len(detail.columns) + 1)
                    self.assertNotIn("Rapproche exact", detail.iloc[0].tolist())
                    self.assertIn(1250.5, detail.iloc[0].tolist())
                self.assertIn("Rapproche exact", audit.iloc[0].tolist())
        pd.testing.assert_frame_equal(self.detail, original)

    def test_excel_accepts_detail_without_control(self):
        frame = self.detail.drop(columns="statut_rapprochement")
        content = create_excel_export({"rapport_journalier_detail": frame})
        result = pd.read_excel(BytesIO(content), sheet_name="Rapport_Journalier_Detail")
        self.assertEqual(list(result.columns), list(frame.columns))

    @unittest.skipUnless(os.environ.get("MPESA_QA_INPUT_DIR"), "Classeurs locaux optionnels")
    def test_real_transactions_export(self):
        paths = sorted(Path(os.environ["MPESA_QA_INPUT_DIR"]).glob("*Transactions*.xlsx"))
        self.assertTrue(paths)
        tx = prepare_transactions(pd.concat([pd.read_excel(p, engine="calamine") for p in paths]))
        day = tx.created_at.max().normalize() - pd.Timedelta(days=1)
        tx = tx.loc[tx.created_at.ge(day) & tx.created_at.lt(day + pd.Timedelta(days=1))].copy()
        data = MpesaPreparedData(transactions=tx, current_savings=pd.DataFrame(),
                                 fixed_savings=pd.DataFrame(), loans=pd.DataFrame(), load_report=pd.DataFrame())
        daily = build_g2_daily_savings_report(data)
        detail = daily["detail"]
        self.assertFalse(detail.empty)
        self.assertIn("statut_rapprochement", detail.columns)
        content = create_excel_export({"rapport_journalier_detail": detail})
        exported = pd.read_excel(BytesIO(content), sheet_name="Rapport_Journalier_Detail")
        self.assertNotIn("statut_rapprochement", exported.columns)
        self.assertEqual(len(detail), len(exported))
        pd.testing.assert_series_equal(detail.groupby("currency_code").montant.sum(),
                                       exported.groupby("currency_code").montant.sum(),
                                       check_dtype=False, check_index_type=False)
        word = create_g2_dat_word(
            {"rapport_journalier_detail": detail, "rapport_journalier_pivot": daily["pivot"],
             "analysis_source_label": "Solution Numérique"},
            period_text=str(day.date()), direction_label="Tous", generated_at=day,
        )
        document = Document(BytesIO(word))
        headers = [c.text for table in document.tables for c in table.rows[0].cells]
        self.assertNotIn("Controle Turbo/G2", headers)
        self.assertNotIn("statut_rapprochement", headers)
        self.assertIn("receipt_no", headers)
        print(f"\nExport reel {day.date()}: {len(detail)} operations, lignes et montants par devise preserves.")
